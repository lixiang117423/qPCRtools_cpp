#!/usr/bin/env python3
"""检查 qPCRtools 投稿包的文字、分页、图片边界和版本声明。"""

from __future__ import annotations

import sys
from pathlib import Path

import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
MS = ROOT / "manuscript"
RENDER = Path("/Users/lixiang/tmp/qpcr_docx_render_v8")

FINAL = {
    "main_pdf": MS / "main_v18_20261029.pdf",
    "main_docx": MS / "main_v18_20261029.docx",
    "supp_pdf": MS / "supplementary_v9_20261029.pdf",
    "supp_docx": MS / "supplementary_v9_20261029.docx",
    "letter_docx": MS / "cover_letter_v7_20261029.docx",
    "letter_pdf": RENDER / "cover_letter_v7_20261029.pdf",
    "main_rendered": RENDER / "main_v18_20261029.pdf",
    "supp_rendered": RENDER / "supplementary_v9_20261029.pdf",
}

BANNED = (
    "[[",
    "todo",
    "tbd",
    "validated equivalence",
    "validated behaviour",
    "both authors read and approved",
    "source for version 1.3.1 is",
    "δct-scale option",
)


def fail(message: str) -> None:
    print(f"FAIL: {message}")
    sys.exit(1)


def check_pdf(path: Path, *, expected_max_pages: int, required: tuple[str, ...] = ()) -> pymupdf.Document:
    if not path.exists() or path.stat().st_size == 0:
        fail(f"missing or empty file: {path}")
    doc = pymupdf.open(path)
    if doc.page_count == 0 or doc.page_count > expected_max_pages:
        fail(f"unexpected page count for {path.name}: {doc.page_count}")
    for page in doc:
        width, height = page.rect.width, page.rect.height
        if abs(width - 595.3) > 1 or abs(height - 841.9) > 1:
            fail(f"non-A4 page in {path.name}: {width:.1f} x {height:.1f}")
    text = "\n".join(page.get_text() for page in doc)
    lowered = text.lower()
    for term in BANNED:
        if term in lowered:
            fail(f"banned text {term!r} in {path.name}")
    for term in required:
        if term.lower() not in lowered:
            fail(f"required text {term!r} missing from {path.name}")
    return doc


def image_boxes(doc: pymupdf.Document) -> list[tuple[int, pymupdf.Rect]]:
    boxes = []
    for page_number, page in enumerate(doc, start=1):
        for image in page.get_images(full=True):
            box = page.get_image_bbox(image)
            if box.width > 2 and box.height > 2:
                boxes.append((page_number, box))
    return boxes


def check_centered_images(doc: pymupdf.Document, expected: int) -> None:
    boxes = image_boxes(doc)
    if len(boxes) != expected:
        fail(f"expected {expected} embedded figures, found {len(boxes)}")
    for page_number, box in boxes:
        page = doc[page_number - 1]
        left = box.x0
        right = page.rect.width - box.x1
        if abs(left - right) > 0.2:
            fail(f"figure is not centered on page {page_number}: L={left:.2f}, R={right:.2f}")
        if box.x0 < 0 or box.y0 < 0 or box.x1 > page.rect.width or box.y1 > page.rect.height:
            fail(f"figure extends outside page {page_number}: {box}")


def check_docx(path: Path, *, images: int, tables: int) -> None:
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    lowered = text.lower()
    for term in BANNED:
        if term in lowered:
            fail(f"banned text {term!r} in {path.name}")
    if len(doc.inline_shapes) != images:
        fail(f"unexpected image count in {path.name}: {len(doc.inline_shapes)}")
    if len(doc.tables) != tables:
        fail(f"unexpected table count in {path.name}: {len(doc.tables)}")


def main() -> None:
    main_latex = check_pdf(
        FINAL["main_pdf"],
        expected_max_pages=10,
        required=("EasyqpcR", "1.3.1", "Holm", "qBase+", "Table S3"),
    )
    supp_latex = check_pdf(
        FINAL["supp_pdf"],
        expected_max_pages=5,
        required=("Table S1", "Table S2", "Table S3", "1.9"),
    )
    main_rendered = check_pdf(FINAL["main_rendered"], expected_max_pages=11)
    supp_rendered = check_pdf(FINAL["supp_rendered"], expected_max_pages=5)
    letter = check_pdf(FINAL["letter_pdf"], expected_max_pages=1)

    check_centered_images(main_rendered, expected=1)
    check_centered_images(supp_rendered, expected=2)
    if image_boxes(letter):
        fail("cover letter unexpectedly contains embedded images")

    first_page = main_rendered[0].get_text()
    for term in ("qPCRtools", "Xiang Li", "Huichuan Huang", "Abstract"):
        if term not in first_page:
            fail(f"Word title page is missing {term!r}")

    table_pages = {
        term: {
            page_number
            for page_number, page in enumerate(main_rendered, start=1)
            if term in page.get_text()
        }
        for term in ("This application", "REST", "LinRegPCR", "qBase", "Web apps")
    }
    common_pages = set.intersection(*table_pages.values())
    if len(common_pages) != 1:
        fail(f"Word Table 1 is split across pages: {table_pages}")

    check_docx(FINAL["main_docx"], images=1, tables=1)
    check_docx(FINAL["supp_docx"], images=2, tables=3)
    check_docx(FINAL["letter_docx"], images=0, tables=0)

    print(
        "OK: "
        f"LaTeX {main_latex.page_count}+{supp_latex.page_count} pages; "
        f"Word {main_rendered.page_count}+{supp_rendered.page_count} pages; "
        f"cover {letter.page_count} page"
    )


if __name__ == "__main__":
    main()
